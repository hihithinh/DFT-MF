import cv2
import math
import os
import argparse
import openpyxl
from timeit import default_timer as timer
import time
import logging
from datetime import datetime

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Extract frames from videos')
    parser.add_argument('dataset', nargs='?', default='CelebDF', choices=['UADFV', 'CelebDF'],
                        help='Dataset name: UADFV or CelebDF (default: CelebDF)')
    parser.add_argument('--video-dir', type=str, help='Directory containing videos (optional)')
    args = parser.parse_args()
    
    dataset_name = args.dataset
    
    logging.basicConfig(
        level=logging.INFO,
        format='[%(asctime)s] %(message)s',
        datefmt='%H:%M:%S'
    )
    log = logging.getLogger(__name__)
    
    log.info("="*60)
    log.info("FRAME EXTRACTION SCRIPT")
    log.info("="*60)
    log.info(f"Dataset: {dataset_name}")
    
    if args.video_dir:
        video_dir = args.video_dir
    else:
        if dataset_name == "UADFV":
            # UADFV has fake/ and real/ folders with videos
            video_dir = os.path.join(BASE_DIR, "dataset", dataset_name)
        else:
            # CelebDF uses Videos/ folder
            video_dir = os.path.join(BASE_DIR, dataset_name, "Videos")
    
    extract_dir = os.path.join(BASE_DIR, dataset_name, "ExtractFrams")
    result_file = os.path.join(BASE_DIR, dataset_name, "Result.xlsx")
    
    log.info(f"Video directory: {video_dir}")
    log.info(f"Output directory: {extract_dir}")
    log.info(f"Results file: {result_file}")
    
    if not os.path.exists(video_dir):
        log.error(f"Video directory not found: {video_dir}")
        log.info(f"Please create the directory and place videos there")
        exit(1)
    
    os.makedirs(extract_dir, exist_ok=True)
    
    # Handle different video directory structures
    if dataset_name == "UADFV" and os.path.exists(os.path.join(video_dir, "fake")) and os.path.exists(os.path.join(video_dir, "real")):
        # UADFV structure: fake/ and real/ folders
        fake_dir = os.path.join(video_dir, "fake")
        real_dir = os.path.join(video_dir, "real")
        
        fake_videos = [os.path.join("fake", f) for f in os.listdir(fake_dir) 
                      if f.lower().endswith(('.mp4', '.avi', '.mov', '.mkv'))]
        real_videos = [os.path.join("real", f) for f in os.listdir(real_dir) 
                      if f.lower().endswith(('.mp4', '.avi', '.mov', '.mkv'))]
        
        video_files = fake_videos + real_videos
        video_files = sorted(video_files)
    else:
        # Standard structure: Videos/ folder
        video_files = [f for f in os.listdir(video_dir) 
                       if f.lower().endswith(('.mp4', '.avi', '.mov', '.mkv'))]
        video_files = sorted(video_files)
    log.info(f"Found {len(video_files)} video(s) to process")
    log.info("="*60)
    
    if os.path.exists(result_file):
        book = openpyxl.load_workbook(result_file)
        sheet = book.active
        sheetCount = 1
        for row in sheet.iter_rows():
            if row[0].value and 'All Frames' in str(row[0].value):
                sheetCount += 1
    else:
        book = openpyxl.Workbook()
        sheet = book.active
        sheetCount = 1
    
    total_start = timer()
    total_frames = 0
    
    for idx, video_file in enumerate(video_files, 1):
        start = timer()
        
        # Handle UADFV subfolder structure (fake/video.mp4, real/video.mp4)
        if os.path.sep in video_file and dataset_name == "UADFV":
            # UADFV: fake/video.mp4 or real/video.mp4
            subfolder, filename = video_file.split(os.path.sep, 1)
            video_name = f"{subfolder}_{os.path.splitext(filename)[0]}"
            video_path = os.path.join(video_dir, video_file)
        else:
            # Standard structure
            video_name = os.path.splitext(video_file)[0]
            video_path = os.path.join(video_dir, video_file)
        
        output_folder = os.path.join(extract_dir, video_name)
        
        os.makedirs(output_folder, exist_ok=True)
        
        log.info(f"[{idx}/{len(video_files)}] Processing: {video_file}")
        
        cap = cv2.VideoCapture(video_path)
        
        if not cap.isOpened():
            log.error(f"  Failed to open video: {video_file}")
            continue
        
        frame_rate = cap.get(cv2.CAP_PROP_FPS)
        total_video_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        
        log.info(f"  FPS: {frame_rate:.2f}, Total frames: {total_video_frames}")
        
        frame_count = 0
        saved_count = 0
        
        while cap.isOpened():
            frameId = cap.get(cv2.CAP_PROP_POS_FRAMES)
            ret, frame = cap.read()
            
            if not ret:
                break
            
            if frameId % 1 == 0:
                filename = os.path.join(output_folder, f"image_{int(frameId)}.jpg")
                cv2.imwrite(filename, frame)
                saved_count += 1
            
            frame_count = int(frameId)
        
        cap.release()
        
        end = timer()
        elapsed = int(end - start)
        
        total_frames += saved_count
        
        log.info(f"  Extracted {saved_count} frames in {elapsed}s")
        log.info(f"  Output: {output_folder}")
        
        sheet[f'A{sheetCount}'] = f"All Frames {video_name}"
        sheet[f'B{sheetCount}'] = str(saved_count)
        sheet[f'K{sheetCount}'] = elapsed
        
        book.save(result_file)
        sheetCount += 1
    
    total_end = timer()
    total_elapsed = int(total_end - total_start)
    
    log.info("="*60)
    log.info("EXTRACTION COMPLETED!")
    log.info(f"Total videos processed: {len(video_files)}")
    log.info(f"Total frames extracted: {total_frames}")
    log.info(f"Total time: {total_elapsed}s ({total_elapsed/60:.1f} minutes)")
    if total_elapsed > 0:
        log.info(f"Average speed: {total_frames/total_elapsed:.1f} frames/second")
    log.info(f"Results saved to: {result_file}")
    log.info(f"Frames saved to: {extract_dir}")
    log.info("="*60)
